# -*- coding: utf-8 -*-
"""
================================================================================
 Criar-PPTX-Reuniao.py  -  Glowy Life Nutrition
 Monta uma apresentacao PowerPoint (.pptx) a partir das imagens da pasta
 da Reuniao de Lideres, uma imagem por slide, na sequencia natural dos nomes.

 NAO precisa ter o Microsoft Office instalado.

 COMO USAR (Windows):
   1. Instale o Python em https://python.org  (marque "Add Python to PATH")
   2. Abra o Prompt de Comando e rode:
          pip install python-pptx
   3. Rode o script:
          python Criar-PPTX-Reuniao.py

   O script procura sozinho a pasta dentro de Downloads.
   Para apontar outra pasta:
          python Criar-PPTX-Reuniao.py "C:\\caminho\\da\\pasta"
================================================================================
"""

import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.dml.color import RGBColor
except ImportError:
    print()
    print("  ERRO: a biblioteca python-pptx nao esta instalada.")
    print()
    print("  Abra o Prompt de Comando e rode:")
    print("      pip install python-pptx")
    print()
    sys.exit(1)


# ------------------------------------------------------------- Configuracoes
NOME_ARQUIVO = "Reuniao_de_Lideres.pptx"

FUNDO = "preto"          # "preto" ou "branco"
PREENCHER_TELA = False   # True = cobre o slide inteiro cortando sobras
                         # False = mostra a imagem inteira, centralizada

# python-pptx nao le .webp. Se houver webp na pasta, o script avisa.
EXTENSOES = (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff")

# Padrao do nome da pasta. Escrito sem acento de proposito: o "ao" de
# "reuniao" e o "i" de "lideres" podem vir acentuados ou nao.
PADRAO_PASTA = re.compile(r"^reuni.*deres", re.IGNORECASE)

# Slide widescreen 16:9 -> 13,333 x 7,5 polegadas
LARGURA_SLIDE = Emu(12192000)
ALTURA_SLIDE = Emu(6858000)


def chave_ordenacao(nome):
    """Ordenacao natural: slide2 vem antes de slide10 (igual ao Windows)."""
    partes = re.split(r"(\d+)", nome.lower())
    return [int(p) if p.isdigit() else p for p in partes]


def descobrir_pasta():
    """Procura a pasta da reuniao dentro de Downloads."""
    downloads = os.path.join(os.path.expanduser("~"), "Downloads")

    if not os.path.isdir(downloads):
        print("  ERRO: nao encontrei a pasta Downloads em: %s" % downloads)
        print('  Rode informando o caminho: python Criar-PPTX-Reuniao.py "C:\\sua\\pasta"')
        return None

    candidatos = [
        os.path.join(downloads, nome)
        for nome in os.listdir(downloads)
        if os.path.isdir(os.path.join(downloads, nome)) and PADRAO_PASTA.match(nome)
    ]

    if not candidatos:
        print("  ERRO: nao encontrei nenhuma pasta de reuniao de lideres em:")
        print("        %s" % downloads)
        print()
        print("  Pastas disponiveis ali:")
        for nome in sorted(os.listdir(downloads)):
            if os.path.isdir(os.path.join(downloads, nome)):
                print("      %s" % nome)
        print()
        print('  Rode informando o caminho: python Criar-PPTX-Reuniao.py "C:\\sua\\pasta"')
        return None

    if len(candidatos) > 1:
        print("  ERRO: encontrei mais de uma pasta possivel:")
        for caminho in candidatos:
            print("      %s" % caminho)
        print()
        print("  Escolha uma e passe o caminho completo como argumento.")
        return None

    return candidatos[0]


def main():
    print()
    print("=============================================")
    print(" GLOWY LIFE - Gerador de PowerPoint")
    print(" Reuniao de Lideres")
    print("=============================================")
    print()

    # ---------------------------------------------- Descobre / valida a pasta
    if len(sys.argv) > 1:
        pasta = sys.argv[1]
    else:
        pasta = descobrir_pasta()
        if pasta is None:
            return 1

    if not os.path.isdir(pasta):
        print("  ERRO: a pasta nao foi encontrada:")
        print("        %s" % pasta)
        return 1

    print("  Pasta: %s" % pasta)

    # ------------------------------------------------- Coleta e ordena imagens
    arquivos = [
        f for f in os.listdir(pasta)
        if os.path.isfile(os.path.join(pasta, f))
    ]

    imagens = [f for f in arquivos if f.lower().endswith(EXTENSOES)]
    imagens.sort(key=chave_ordenacao)

    webp = [f for f in arquivos if f.lower().endswith(".webp")]
    if webp:
        print()
        print("  AVISO: %d arquivo(s) .webp serao IGNORADOS." % len(webp))
        print("         A biblioteca python-pptx nao le esse formato.")
        print("         Use o script Criar-PPTX-Reuniao.ps1 (PowerShell), que aceita webp,")
        print("         ou converta as imagens para .png antes de rodar.")

    if not imagens:
        print("  ERRO: nenhuma imagem encontrada na pasta.")
        print("  Formatos aceitos: %s" % ", ".join(EXTENSOES))
        return 1

    print()
    print("  OK  %d imagens encontradas. Ordem dos slides:" % len(imagens))
    for i, nome in enumerate(imagens, start=1):
        print("      %3d. %s" % (i, nome))
    print()
    print("  Confira a ordem acima.")
    try:
        resposta = input("  Digite S para montar a apresentacao (ou ENTER para cancelar): ")
    except EOFError:
        resposta = "s"
    if not resposta.strip().lower().startswith("s"):
        print("  Cancelado. Renomeie os arquivos com prefixo numerico (01_, 02_...) e rode de novo.")
        return 0
    print()

    # ------------------------------------------------------ Monta a apresentacao
    print("  Montando os slides...")

    prs = Presentation()
    prs.slide_width = LARGURA_SLIDE
    prs.slide_height = ALTURA_SLIDE

    layout_em_branco = prs.slide_layouts[6]
    cor_fundo = RGBColor(0xFF, 0xFF, 0xFF) if FUNDO == "branco" else RGBColor(0x00, 0x00, 0x00)

    for i, nome in enumerate(imagens, start=1):
        caminho = os.path.join(pasta, nome)
        slide = prs.slides.add_slide(layout_em_branco)

        # Fundo solido, para que as bordas fiquem com acabamento intencional
        fundo = slide.background.fill
        fundo.solid()
        fundo.fore_color.rgb = cor_fundo

        # Insere a imagem no tamanho original e depois escala
        figura = slide.shapes.add_picture(caminho, 0, 0)

        if PREENCHER_TELA:
            # Cobre o slide inteiro (pode cortar as sobras)
            escala = max(
                LARGURA_SLIDE / figura.width,
                ALTURA_SLIDE / figura.height,
            )
        else:
            # Mostra a imagem inteira dentro do slide
            escala = min(
                LARGURA_SLIDE / figura.width,
                ALTURA_SLIDE / figura.height,
            )

        figura.width = int(figura.width * escala)
        figura.height = int(figura.height * escala)

        # Centraliza
        figura.left = int((LARGURA_SLIDE - figura.width) / 2)
        figura.top = int((ALTURA_SLIDE - figura.height) / 2)

        print("      slide %3d  <-  %s" % (i, nome))

    # ------------------------------------------------------------- Salva tudo
    caminho_saida = os.path.join(pasta, NOME_ARQUIVO)
    prs.save(caminho_saida)

    print()
    print("  OK  Apresentacao criada com %d slides:" % len(imagens))
    print("        %s" % caminho_saida)
    print()
    return 0


if __name__ == "__main__":
    codigo = main()
    if os.name == "nt":
        try:
            input("Pressione ENTER para fechar")
        except EOFError:
            pass
    sys.exit(codigo)
