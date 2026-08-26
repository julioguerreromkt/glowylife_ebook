# -*- coding: utf-8 -*-
"""
================================================================================
 Criar-PPTX.py  -  Glowy Life Nutrition
 Monta uma apresentacao PowerPoint (.pptx) a partir das imagens de uma pasta,
 uma imagem por slide, na sequencia natural dos nomes dos arquivos.

 NAO precisa ter o Microsoft Office instalado.

 COMO USAR (Windows):
   1. Instale o Python em https://python.org  (marque "Add Python to PATH")
   2. Abra o Prompt de Comando e rode:
          pip install python-pptx
   3. Rode o script:
          python Criar-PPTX.py

   Para usar outra pasta:
          python Criar-PPTX.py "C:\\caminho\\da\\pasta"
================================================================================
"""

import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
except ImportError:
    print()
    print("  ERRO: a biblioteca python-pptx nao esta instalada.")
    print()
    print("  Abra o Prompt de Comando e rode:")
    print("      pip install python-pptx")
    print()
    sys.exit(1)


# ------------------------------------------------------------- Configuracoes
PASTA_PADRAO = r"C:\Users\glowy\Downloads\Live_programaNeurobalance"
NOME_ARQUIVO = "Live_Programa_Neurobalance.pptx"

FUNDO = "preto"          # "preto" ou "branco"
PREENCHER_TELA = False   # True = cobre o slide inteiro cortando sobras
                         # False = mostra a imagem inteira, centralizada

EXTENSOES = (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff")

# Slide widescreen 16:9 -> 13,333 x 7,5 polegadas
LARGURA_SLIDE = Emu(12192000)
ALTURA_SLIDE = Emu(6858000)


def chave_ordenacao(nome):
    """Ordenacao natural: slide2 vem antes de slide10 (igual ao Windows)."""
    partes = re.split(r"(\d+)", nome.lower())
    return [int(p) if p.isdigit() else p for p in partes]


def main():
    pasta = sys.argv[1] if len(sys.argv) > 1 else PASTA_PADRAO

    print()
    print("=============================================")
    print(" GLOWY LIFE - Gerador de PowerPoint")
    print("=============================================")
    print()

    # ------------------------------------------------------------ Validacao
    if not os.path.isdir(pasta):
        print("  ERRO: a pasta nao foi encontrada:")
        print("        %s" % pasta)
        print()
        print("  Rode informando o caminho correto, por exemplo:")
        print('      python Criar-PPTX.py "C:\\sua\\pasta"')
        print()
        return 1

    print("  Pasta: %s" % pasta)

    # ------------------------------------------------- Coleta e ordena imagens
    imagens = [
        f for f in os.listdir(pasta)
        if os.path.isfile(os.path.join(pasta, f))
        and f.lower().endswith(EXTENSOES)
    ]
    imagens.sort(key=chave_ordenacao)

    if not imagens:
        print("  ERRO: nenhuma imagem encontrada na pasta.")
        print("  Formatos aceitos: %s" % ", ".join(EXTENSOES))
        print()
        return 1

    print("  OK  %d imagens encontradas. Ordem dos slides:" % len(imagens))
    for i, nome in enumerate(imagens, start=1):
        print("      %3d. %s" % (i, nome))
    print()

    # ------------------------------------------------------ Monta a apresentacao
    print("  Montando os slides...")

    prs = Presentation()
    prs.slide_width = LARGURA_SLIDE
    prs.slide_height = ALTURA_SLIDE

    layout_em_branco = prs.slide_layouts[6]
    cor_fundo = RGBColor(0xFF, 0xFF, 0xFF) if FUNDO == "branco" else RGBColor(0x00, 0x00, 0x00)

    for i, nome in enumerate(imagens, start=1):
        caminho = os.path.join(pasta, nome)
        slide = prs.slides.add_slide(layout_em_branco)

        # Fundo solido, para que as bordas fiquem com acabamento intencional
        fundo = slide.background.fill
        fundo.solid()
        fundo.fore_color.rgb = cor_fundo

        # Insere a imagem no tamanho original e depois escala
        figura = slide.shapes.add_picture(caminho, 0, 0)

        if PREENCHER_TELA:
            # Cobre o slide inteiro (pode cortar as sobras)
            escala = max(
                LARGURA_SLIDE / figura.width,
                ALTURA_SLIDE / figura.height,
            )
        else:
            # Mostra a imagem inteira dentro do slide
            escala = min(
                LARGURA_SLIDE / figura.width,
                ALTURA_SLIDE / figura.height,
            )

        figura.width = int(figura.width * escala)
        figura.height = int(figura.height * escala)

        # Centraliza
        figura.left = int((LARGURA_SLIDE - figura.width) / 2)
        figura.top = int((ALTURA_SLIDE - figura.height) / 2)

        print("      slide %3d  <-  %s" % (i, nome))

    # ------------------------------------------------------------- Salva tudo
    caminho_saida = os.path.join(pasta, NOME_ARQUIVO)
    prs.save(caminho_saida)

    print()
    print("  OK  Apresentacao criada com %d slides:" % len(imagens))
    print("        %s" % caminho_saida)
    print()
    return 0


if __name__ == "__main__":
    codigo = main()
    if os.name == "nt":
        try:
            input("Pressione ENTER para fechar")
        except EOFError:
            pass
    sys.exit(codigo)
